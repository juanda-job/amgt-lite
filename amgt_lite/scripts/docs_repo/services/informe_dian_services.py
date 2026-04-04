import pandas as pd
import unicodedata
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from pathlib import Path

def remove_accents(text: str) -> str:
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')

def normalize_text(x):
    if pd.isna(x):
        return x
    x = str(x).strip().upper()
    x = remove_accents(x)
    x = x.lstrip('-')
    return x

def generarInforme(archivo):
    try:
        # Detectar tipo de archivo por nombre
        nombre = getattr(archivo, "name", str(archivo))
        if nombre.endswith(".csv"):
            df = pd.read_csv(archivo, sep=";", encoding="latin1")
        elif nombre.endswith((".xlsx", ".xls")):
            df = pd.read_excel(archivo, engine="openpyxl")
        else:
            raise ValueError(f"Formato no soportado: {Path(nombre).suffix}")



        for col in ["IVA", "Total"]:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

        if {"Tipo de documento", "Grupo"}.issubset(df.columns):
            resumen = (
                df.groupby(["Tipo de documento", "Grupo"], as_index=False)[["IVA", "Total"]]
                  .sum()
            )
            resumen["Antes de IVA"] = resumen["Total"] - resumen["IVA"]
            resumen = resumen[["Tipo de documento", "Grupo", "Antes de IVA", "IVA", "Total"]]
            resumen = resumen.sort_values(by="Grupo")

            bloques = []
            for grupo, datos in resumen.groupby("Grupo"):
                encabezados = pd.DataFrame(
                    [["Tipo de documento", "Grupo", "Antes de IVA", "IVA", "Total"]],
                    columns=resumen.columns
                )
                bloques.append(encabezados)
                bloques.append(datos)
                bloques.append(pd.DataFrame([["", "", "", "", ""]] * 3, columns=resumen.columns))
            resumen_final = pd.concat(bloques, ignore_index=True)
        else:
            resumen_final = pd.DataFrame()

        # 👉 Exportar a buffer en memoria
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="datos", index=False)
            resumen_final.to_excel(writer, sheet_name="resumen", index=False, header=False)

        # 👉 Aplicar bordes con openpyxl
        output.seek(0)
        wb = load_workbook(output)
        ws = wb["resumen"]

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    cell.border = thin_border

        # Guardar nuevamente en buffer
        final_output = BytesIO()
        wb.save(final_output)
        final_output.seek(0)

        # 👉 Devolver el buffer final (para que el view lo use)
        return final_output

    except Exception as e:
        raise ValueError(f"Error al procesar datos: {str(e)}")